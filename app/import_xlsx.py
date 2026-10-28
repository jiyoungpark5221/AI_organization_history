"""환경부 Pilot 엑셀 가져오기 (참고 자료 입력·매핑 시험용).

- 엑셀은 read_only 로만 열고 수정하지 않는다 (가져오기 전후 sha256 을 보고서에 기록).
- org / org_identifier(COMMON, GCC) 와 ref_* 참고 테이블에만 쓴다. 확정 데이터·검토 이력·후보,
  collection_run / gcc_record / change_event 는 쓰지 않는다 (이 파일은 일자별 GCC 스냅샷이 아니다).
- 파일의 상태값(REVIEW_STATUS, MASTER_STATUS 등)은 원본 그대로 *_src 컬럼에 보존할 뿐, 우리 시스템의 승인이 아니다.
- OG 전거ID 는 이 파일에서 확보된 것이 없으므로 org_identifier 'OG' 에 넣지 않는다.

사용: python -m app.import_xlsx <엑셀경로> [--db 경로]
"""
import argparse
import hashlib
import json
import re
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import openpyxl

from .db import connect, init_db, now, tx

SHEET_ORG = "01_ORG_MASTER"
SHEET_SRC = "02_SOURCE_MAP"
SHEET_REL = "03_ORG_RELATION"
SHEET_WORK = "09_ORG_WORK"
USED_SHEETS = [SHEET_ORG, SHEET_SRC, SHEET_REL, SHEET_WORK]

_NULLISH = {"", "NULL", "None", "none", "null", "nan"}
_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}")
_NO_PARENT = {"0000000"}


class ImportError_(Exception):
    pass


def sha256_of(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def clean(v):
    """빈 값·'NULL'·'None' → None, 숫자형 코드는 정수 문자열, 나머지는 공백 제거한 문자열."""
    if v is None:
        return None
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return None if s in _NULLISH else s


def to_date(v, stats: Counter | None = None):
    """날짜 → 'YYYY-MM-DD'. 날짜로 읽을 수 없는 값은 원문 그대로 두고 stats['nonstandard'] 로 센다."""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    s = clean(v)
    if s is None or s == "미확인":
        return None
    if _DATE.match(s):
        return s[:10]
    if stats is not None:
        stats["nonstandard"] += 1
    return s


def _read_sheet(wb, name: str, first_col: str) -> list[dict]:
    """헤더 행(첫 열 이름이 first_col)을 찾아 그 아래 행을 dict 목록으로 반환."""
    if name not in wb.sheetnames:
        raise ImportError_(f"시트 없음: {name}")
    rows = wb[name].iter_rows(values_only=True)
    header = None
    out = []
    for i, r in enumerate(rows):
        if header is None:
            if i > 15:
                break
            if r and clean(r[0]) == first_col:
                header = [clean(c) for c in r]
            continue
        if any(clean(c) is not None for c in r):
            out.append({h: c for h, c in zip(header, r) if h})
    if header is None:
        raise ImportError_(f"{name}: 헤더 행({first_col})을 찾지 못함")
    return out


def read_workbook(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return {
            "sheetnames": list(wb.sheetnames),
            "org": _read_sheet(wb, SHEET_ORG, "COMMON_ORG_ID"),
            "src": _read_sheet(wb, SHEET_SRC, "SOURCE_ROW_ID"),
            "rel": _read_sheet(wb, SHEET_REL, "RELATION_ID"),
            "work": _read_sheet(wb, SHEET_WORK, "WORK_ID"),
        }
    finally:
        wb.close()


def _codes(v) -> list[str]:
    s = clean(v)
    return [c for c in re.split(r"[,\s;/]+", s) if c] if s else []


# ── 스키마 보정 ────────────────────────────────────────────────

def ensure_schema(conn) -> None:
    """새 DB 는 schema.sql 그대로. 이전에 만든 DB 의 org_identifier CHECK 에 COMMON 이 없으면 표를 다시 만든다."""
    row = conn.execute("SELECT sql FROM sqlite_master WHERE name='org_identifier'").fetchone()
    if row and "'COMMON'" not in row["sql"]:
        conn.execute("PRAGMA foreign_keys=OFF")
        try:
            conn.execute("BEGIN IMMEDIATE")
            conn.execute("ALTER TABLE org_identifier RENAME TO org_identifier_old")
            conn.execute(
                "CREATE TABLE org_identifier (org_id INTEGER NOT NULL REFERENCES org(org_id), "
                "id_type TEXT NOT NULL CHECK (id_type IN ('GCC','OG','COMMON')), "
                "id_value TEXT NOT NULL, PRIMARY KEY (id_type, id_value))"
            )
            conn.execute("INSERT INTO org_identifier SELECT org_id, id_type, id_value FROM org_identifier_old")
            conn.execute("DROP TABLE org_identifier_old")
            conn.execute("COMMIT")
        except BaseException:
            conn.execute("ROLLBACK")
            raise
        finally:
            conn.execute("PRAGMA foreign_keys=ON")
    init_db(conn)


# ── 쓰기 단계 (모두 하나의 트랜잭션 안에서 호출) ────────────────

def _write_orgs(conn, data: dict, report: dict) -> dict[str, int]:
    """공통기관ID → org_id. org 1개당 COMMON 식별자 1건 + 02 의 GCC 코드들."""
    names: dict[str, str] = {}
    for r in data["org"]:
        cid = clean(r.get("COMMON_ORG_ID"))
        if cid:
            names[cid] = clean(r.get("DISPLAY_NAME")) or cid
    gcc_by_common: dict[str, list[str]] = {}
    for r in data["src"]:
        cid = clean(r.get("COMMON_ORG_ID"))
        if not cid:
            continue
        if cid not in names:
            names[cid] = clean(r.get("SOURCE_FULL_ORG_NAME")) or cid
            report["counts"]["orgs_only_in_source_map"] += 1
        if clean(r.get("SOURCE_SYSTEM")) == "GCC" and clean(r.get("SOURCE_CODE")):
            codes = gcc_by_common.setdefault(cid, [])
            code = clean(r["SOURCE_CODE"])
            if code not in codes:
                codes.append(code)

    org_of: dict[str, int] = {}
    c = report["counts"]
    for cid, name in names.items():
        row = conn.execute(
            "SELECT org_id FROM org_identifier WHERE id_type='COMMON' AND id_value=?", (cid,)
        ).fetchone()
        org_id = row["org_id"] if row else None
        if org_id is None:  # 이미 같은 GCC 코드로 만들어진 기관(예: 승인으로 생성)이 있으면 재사용
            for code in gcc_by_common.get(cid, []):
                row = conn.execute(
                    "SELECT org_id FROM org_identifier WHERE id_type='GCC' AND id_value=?", (code,)
                ).fetchone()
                if row:
                    org_id = row["org_id"]
                    c["orgs_matched_existing"] += 1
                    break
        if org_id is None:
            org_id = conn.execute("INSERT INTO org(org_name) VALUES(?)", (name,)).lastrowid
            c["orgs_created"] += 1
        org_of[cid] = org_id
        cur = conn.execute(
            "INSERT OR IGNORE INTO org_identifier(org_id, id_type, id_value) VALUES(?, 'COMMON', ?)",
            (org_id, cid),
        )
        c["common_identifiers_added"] += cur.rowcount
        for code in gcc_by_common.get(cid, []):
            owner = conn.execute(
                "SELECT org_id FROM org_identifier WHERE id_type='GCC' AND id_value=?", (code,)
            ).fetchone()
            if owner is None:
                conn.execute(
                    "INSERT INTO org_identifier(org_id, id_type, id_value) VALUES(?, 'GCC', ?)", (org_id, code)
                )
                c["gcc_identifiers_added"] += 1
            elif owner["org_id"] != org_id:
                report["unlinked"]["gcc_code_conflicts"].append(
                    {"gcc_code": code, "common_org_id": cid, "kept_org_id": owner["org_id"]}
                )
    return org_of


def _write_records(conn, data, org_of, import_id, report) -> None:
    master = {clean(r.get("COMMON_ORG_ID")): r for r in data["org"]}
    dstats = report["date_stats"]
    rows = []
    for r in data["src"]:
        cid = clean(r.get("COMMON_ORG_ID"))
        m = master.get(cid, {})
        rows.append((
            clean(r.get("SOURCE_ROW_ID")), import_id, org_of.get(cid), cid,
            clean(r.get("SOURCE_SYSTEM")) or "", clean(r.get("SOURCE_CODE")),
            clean(r.get("SOURCE_FULL_ORG_NAME")), clean(r.get("SOURCE_LOWEST_ORG_NAME")),
            clean(r.get("DEPTH")), clean(r.get("PARENT_CODE")), clean(r.get("TOP_CODE")),
            clean(r.get("PREVIOUS_CODE")),
            to_date(r.get("SOURCE_START"), dstats), to_date(r.get("SOURCE_END"), dstats),
            clean(r.get("SOURCE_STATUS")), clean(m.get("MASTER_STATUS")), clean(m.get("PERIOD_STATUS")),
            clean(r.get("SOURCE_FILE")), clean(r.get("SOURCE_SHEET")),
            clean(r.get("RAW_SOURCE_ID")), clean(r.get("RAW_SOURCE_LOCATOR")),
        ))
    conn.executemany(
        "INSERT INTO ref_org_record(source_row_id, import_id, org_id, common_org_id, source_system, source_code,"
        " full_name, lowest_name, depth, parent_code, top_code, previous_code, source_start, source_end,"
        " source_status, master_status_src, period_status_src, source_file, source_sheet, raw_source_id, raw_locator)"
        " VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)"
        " ON CONFLICT(source_row_id) DO UPDATE SET import_id=excluded.import_id, org_id=excluded.org_id,"
        " common_org_id=excluded.common_org_id, source_system=excluded.source_system, source_code=excluded.source_code,"
        " full_name=excluded.full_name, lowest_name=excluded.lowest_name, depth=excluded.depth,"
        " parent_code=excluded.parent_code, top_code=excluded.top_code, previous_code=excluded.previous_code,"
        " source_start=excluded.source_start, source_end=excluded.source_end, source_status=excluded.source_status,"
        " master_status_src=excluded.master_status_src, period_status_src=excluded.period_status_src,"
        " source_file=excluded.source_file, source_sheet=excluded.source_sheet,"
        " raw_source_id=excluded.raw_source_id, raw_locator=excluded.raw_locator",
        rows,
    )


def _write_works(conn, data, org_of, import_id, report) -> None:
    rows = []
    for r in data["work"]:
        cid = clean(r.get("COMMON_ORG_ID"))
        rows.append((
            clean(r.get("WORK_ID")), import_id, org_of.get(cid), cid, clean(r.get("OG_CODE")),
            clean(r.get("REPRESENTATIVE_NAME")), clean(r.get("RAW_WORK_TEXT")),
            clean(r.get("RAW_START")), clean(r.get("RAW_END")),
            clean(r.get("SOURCE_ID")), clean(r.get("SOURCE_FILE")), clean(r.get("SOURCE_PAGE")),
            clean(r.get("SOURCE_SEQUENCE")), clean(r.get("BRM_FUNCTION_STATUS")),
            to_date(r.get("WORK_VALID_FROM")), to_date(r.get("WORK_VALID_TO")),
            clean(r.get("WORK_PERIOD_STATUS")), clean(r.get("WORK_ADOPTION_STATUS")),
            clean(r.get("LAW_REFERENCE_RAW")),
        ))
    conn.executemany(
        "INSERT INTO ref_org_work(work_id, import_id, org_id, common_org_id, og_code_raw, representative_name,"
        " raw_work_text, raw_start, raw_end, source_id, source_file, source_page, source_sequence,"
        " brm_function_status, work_valid_from, work_valid_to, work_period_status, work_adoption_status,"
        " law_reference_raw) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)"
        " ON CONFLICT(work_id) DO UPDATE SET import_id=excluded.import_id, org_id=excluded.org_id,"
        " common_org_id=excluded.common_org_id, og_code_raw=excluded.og_code_raw,"
        " representative_name=excluded.representative_name, raw_work_text=excluded.raw_work_text,"
        " raw_start=excluded.raw_start, raw_end=excluded.raw_end, source_id=excluded.source_id,"
        " source_file=excluded.source_file, source_page=excluded.source_page,"
        " source_sequence=excluded.source_sequence, brm_function_status=excluded.brm_function_status,"
        " work_valid_from=excluded.work_valid_from, work_valid_to=excluded.work_valid_to,"
        " work_period_status=excluded.work_period_status, work_adoption_status=excluded.work_adoption_status,"
        " law_reference_raw=excluded.law_reference_raw",
        rows,
    )


def _write_relations(conn, data, org_of, import_id, report) -> None:
    dstats = report["date_stats"]
    rows = []
    for r in data["rel"]:
        f, t = clean(r.get("FROM_ORG_ID")), clean(r.get("TO_ORG_ID"))
        rows.append((
            clean(r.get("RELATION_ID")), import_id, clean(r.get("RELATION_LAYER")) or "", clean(r.get("RELATION_TYPE")),
            org_of.get(f), org_of.get(t), f, t,
            clean(r.get("FROM_SOURCE_CODE")), clean(r.get("TO_SOURCE_CODE")),
            clean(r.get("SOURCE_SYSTEM")), clean(r.get("SOURCE_FIELD")),
            to_date(r.get("OBSERVED_FROM"), dstats), to_date(r.get("OBSERVED_TO"), dstats),
            to_date(r.get("RELATION_VALID_FROM"), dstats), to_date(r.get("RELATION_VALID_TO"), dstats),
            clean(r.get("REVIEW_STATUS")), clean(r.get("AI_USE_POLICY")), clean(r.get("EVIDENCE")),
            clean(r.get("RELATION_DATE_STATUS")),
        ))
    conn.executemany(
        "INSERT INTO ref_org_relation(relation_id, import_id, layer, relation_type, from_org_id, to_org_id,"
        " from_common_org_id, to_common_org_id, from_source_code, to_source_code, source_system, source_field,"
        " observed_from, observed_to, relation_valid_from, relation_valid_to, review_status_src,"
        " ai_use_policy_src, evidence, date_status_src) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)"
        " ON CONFLICT(relation_id) DO UPDATE SET import_id=excluded.import_id, layer=excluded.layer,"
        " relation_type=excluded.relation_type, from_org_id=excluded.from_org_id, to_org_id=excluded.to_org_id,"
        " from_common_org_id=excluded.from_common_org_id, to_common_org_id=excluded.to_common_org_id,"
        " from_source_code=excluded.from_source_code, to_source_code=excluded.to_source_code,"
        " source_system=excluded.source_system, source_field=excluded.source_field,"
        " observed_from=excluded.observed_from, observed_to=excluded.observed_to,"
        " relation_valid_from=excluded.relation_valid_from, relation_valid_to=excluded.relation_valid_to,"
        " review_status_src=excluded.review_status_src, ai_use_policy_src=excluded.ai_use_policy_src,"
        " evidence=excluded.evidence, date_status_src=excluded.date_status_src",
        rows,
    )


# ── 보고서 계산 (실제 파일에서 계산, 추정치 하드코딩 없음) ─────────

def _analyze(data: dict, org_of_probe: set[str]) -> dict:
    src, rel, work = data["src"], data["rel"], data["work"]
    gcc_rows = [r for r in src if clean(r.get("SOURCE_SYSTEM")) == "GCC"]
    gcc_codes = {clean(r.get("SOURCE_CODE")) for r in gcc_rows if clean(r.get("SOURCE_CODE"))}

    def missing(values) -> tuple[int, list[str]]:
        bad = [v for v in values if v not in gcc_codes and v not in _NO_PARENT]
        return len(bad), sorted(set(bad))

    parents = [c for r in gcc_rows for c in _codes(r.get("PARENT_CODE"))]
    prevs = [c for r in gcc_rows for c in _codes(r.get("PREVIOUS_CODE"))]
    rel_from = [c for r in rel for c in _codes(r.get("FROM_SOURCE_CODE"))]
    rel_to = [c for r in rel for c in _codes(r.get("TO_SOURCE_CODE"))]
    n_p, d_p = missing(parents)
    n_v, d_v = missing(prevs)
    n_f, d_f = missing(rel_from)
    n_t, d_t = missing(rel_to)
    rel_org_unresolved = sum(
        1 for r in rel
        if clean(r.get("FROM_ORG_ID")) not in org_of_probe or clean(r.get("TO_ORG_ID")) not in org_of_probe
    )

    def blank(rows, col):
        return sum(1 for r in rows if to_date(r.get(col)) is None)

    return {
        "unlinked": {
            "rows_without_code": [
                {"source_row_id": clean(r.get("SOURCE_ROW_ID")), "source_system": clean(r.get("SOURCE_SYSTEM")),
                 "name": clean(r.get("SOURCE_FULL_ORG_NAME"))}
                for r in src if not clean(r.get("SOURCE_CODE"))
            ],
            "parent_code_not_in_gcc": {"rows": n_p, "distinct_codes": len(d_p), "sample": d_p[:20]},
            "previous_code_not_in_gcc": {"rows": n_v, "distinct_codes": len(d_v), "sample": d_v[:20]},
            "relation_from_code_not_in_gcc": {"rows": n_f, "distinct_codes": len(d_f), "sample": d_f[:20]},
            "relation_to_code_not_in_gcc": {"rows": n_t, "distinct_codes": len(d_t), "sample": d_t[:20]},
            "relation_to_code_blank": sum(1 for r in rel if not clean(r.get("TO_SOURCE_CODE"))),
            "relation_org_unresolved": rel_org_unresolved,
            "gcc_code_conflicts": [],
        },
        "unknown_dates": {
            "org_record_total": len(src),
            "source_start_blank": blank(src, "SOURCE_START"),
            "source_end_blank": blank(src, "SOURCE_END"),
            "gcc_source_start_blank": blank(gcc_rows, "SOURCE_START"),
            "gcc_source_end_blank": blank(gcc_rows, "SOURCE_END"),
            "relation_total": len(rel),
            "relation_valid_from_blank": blank(rel, "RELATION_VALID_FROM"),
            "relation_valid_to_blank": blank(rel, "RELATION_VALID_TO"),
            "relation_observed_from_blank": blank(rel, "OBSERVED_FROM"),
            "relation_observed_to_blank": blank(rel, "OBSERVED_TO"),
            "work_total": len(work),
            "work_raw_start_blank": blank(work, "RAW_START"),
            "work_raw_end_blank": blank(work, "RAW_END"),
            "work_valid_from_unknown": blank(work, "WORK_VALID_FROM"),
            "work_valid_to_unknown": blank(work, "WORK_VALID_TO"),
            "work_period_status_unknown": sum(
                1 for r in work if clean(r.get("WORK_PERIOD_STATUS")) in (None, "미확인")
            ),
        },
    }


def run_import(conn, path, *, source_name: str | None = None) -> dict:
    """엑셀을 읽어 참고 테이블에 가져온다 (한 트랜잭션). 보고서 dict 를 반환한다."""
    path = Path(path)
    sha_before = sha256_of(path)
    data = read_workbook(path)
    ensure_schema(conn)

    common_ids = {clean(r.get("COMMON_ORG_ID")) for r in data["org"] if clean(r.get("COMMON_ORG_ID"))}
    common_ids |= {clean(r.get("COMMON_ORG_ID")) for r in data["src"] if clean(r.get("COMMON_ORG_ID"))}
    report = _analyze(data, common_ids)
    report.update({
        "file_name": source_name or path.name,
        "sha256_before": sha_before,
        "counts": Counter(),
        "date_stats": Counter(),
    })

    with tx(conn):
        org_of = _write_orgs(conn, data, report)
        import_id = conn.execute(
            "INSERT INTO import_run(file_name, file_sha256, imported_at, report_json) VALUES(?,?,?,'{}')",
            (report["file_name"], sha_before, now()),
        ).lastrowid
        _write_records(conn, data, org_of, import_id, report)
        _write_works(conn, data, org_of, import_id, report)
        _write_relations(conn, data, org_of, import_id, report)

        c = report["counts"]
        c["orgs_in_file"] = len(org_of)
        c["ref_org_record"] = dict(Counter(clean(r.get("SOURCE_SYSTEM")) or "?" for r in data["src"]))
        c["ref_org_work"] = len(data["work"])
        c["ref_org_relation"] = dict(Counter(clean(r.get("RELATION_LAYER")) or "?" for r in data["rel"]))
        report["counts"] = dict(c)
        report["date_stats"] = dict(report["date_stats"])
        report["og_authority_id"] = {
            "imported": 0,
            "reason": "이 파일에는 OG 전거ID가 기재되어 있지 않다. 09_ORG_WORK 의 OG_CODE 는 같은 기관의 GCC 코드와 "
                      "같은 값이어서 전거ID 로 보지 않고 ref_org_work.og_code_raw 로만 보존했다.",
        }
        report["og_code_equals_gcc_code"] = _og_equals_gcc(data)
        report["not_imported_sheets"] = [s for s in data["sheetnames"] if s not in USED_SHEETS]
        report["change_detection_notice"] = (
            "이 파일은 일자별 GCC 스냅샷이 아니다. 실제 변경 탐지 검증에는 별도의 변경 전·후 GCC 자료가 필요하며 "
            "이 가져오기로 검증된 것이 아니다. collection_run/gcc_record/change_event 는 만들지 않았다."
        )
        # 보고서에 사용 기록(원본 상태 보존·승인 아님) 명시
        report["status_note"] = "파일의 REVIEW_STATUS/MASTER_STATUS 등은 원본 상태로만 보존했으며 우리 시스템의 승인이 아니다."
        sha_after = sha256_of(path)
        report["sha256_after"] = sha_after
        report["source_file_unchanged"] = sha_before == sha_after
        conn.execute("UPDATE import_run SET report_json=? WHERE import_id=?",
                     (json.dumps(report, ensure_ascii=False), import_id))
        report["import_id"] = import_id
    return report


def _og_equals_gcc(data) -> dict:
    gcc: dict[str, set[str]] = {}
    for r in data["src"]:
        if clean(r.get("SOURCE_SYSTEM")) == "GCC" and clean(r.get("SOURCE_CODE")):
            gcc.setdefault(clean(r.get("COMMON_ORG_ID")), set()).add(clean(r["SOURCE_CODE"]))
    with_og = [r for r in data["work"] if clean(r.get("OG_CODE"))]
    same = sum(1 for r in with_og if clean(r["OG_CODE"]) in gcc.get(clean(r.get("COMMON_ORG_ID")), set()))
    return {"work_rows_with_og_code": len(with_og), "same_as_gcc_code_of_same_org": same}


def format_report(rep: dict) -> str:
    c, u, d = rep["counts"], rep["unlinked"], rep["unknown_dates"]
    L = [
        f"가져오기 완료 (import_id={rep.get('import_id')}) — {rep['file_name']}",
        f"원본 sha256 전={rep['sha256_before'][:16]}… 후={rep['sha256_after'][:16]}…  변경 없음={rep['source_file_unchanged']}",
        "",
        "[가져온 건수]",
        f"  기관(org): 파일 {c['orgs_in_file']}건 (신규 {c.get('orgs_created', 0)}, 기존 재사용 {c.get('orgs_matched_existing', 0)})",
        f"  식별자: 공통기관ID {c.get('common_identifiers_added', 0)}건 추가, GCC 코드 {c.get('gcc_identifiers_added', 0)}건 추가, OG 전거ID {rep['og_authority_id']['imported']}건",
        f"  ref_org_record: {c['ref_org_record']}",
        f"  ref_org_work: {c['ref_org_work']}",
        f"  ref_org_relation(layer별): {c['ref_org_relation']}",
        "",
        "[연결하지 못한 코드]",
        f"  코드 없는 원천행: {len(u['rows_without_code'])}건 {[x['source_row_id'] for x in u['rows_without_code']]}",
        f"  상위기관코드가 GCC 집합에 없음: {u['parent_code_not_in_gcc']['rows']}행 / {u['parent_code_not_in_gcc']['distinct_codes']}종 {u['parent_code_not_in_gcc']['sample'][:10]}",
        f"  이전기관코드가 GCC 집합에 없음: {u['previous_code_not_in_gcc']['rows']}행 / {u['previous_code_not_in_gcc']['distinct_codes']}종 {u['previous_code_not_in_gcc']['sample'][:10]}",
        f"  관계 출발코드 없음: {u['relation_from_code_not_in_gcc']['rows']}행, 도착코드 없음: {u['relation_to_code_not_in_gcc']['rows']}행, 도착코드 공란: {u['relation_to_code_blank']}행",
        f"  관계의 기관ID 미해결: {u['relation_org_unresolved']}건, GCC 코드 중복 충돌: {len(u['gcc_code_conflicts'])}건",
        "",
        "[미확인 날짜]",
        f"  원천행 시작일 공란 {d['source_start_blank']}/{d['org_record_total']}, 종료일 공란 {d['source_end_blank']}/{d['org_record_total']}",
        f"  관계 유효기간(from/to) 공란 {d['relation_valid_from_blank']}/{d['relation_valid_to_blank']} (전체 {d['relation_total']}건)",
        f"  업무 RAW 시작/종료 공란 {d['work_raw_start_blank']}/{d['work_raw_end_blank']}, 업무기간 미확인 {d['work_period_status_unknown']}/{d['work_total']}건",
        "",
        f"[OG 전거ID] 0건 — {rep['og_authority_id']['reason']}",
        f"[가져오지 않은 시트] {', '.join(rep['not_imported_sheets'])}",
        f"[주의] {rep['change_detection_notice']}",
        f"[주의] {rep['status_note']}",
    ]
    return "\n".join(L)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="환경부 Pilot 엑셀을 참고 테이블로 가져오기")
    ap.add_argument("xlsx")
    ap.add_argument("--db", default=None, help="DB 경로 (기본: APP_DB 또는 data/app.db)")
    args = ap.parse_args(argv)
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    conn = connect(args.db)
    try:
        rep = run_import(conn, args.xlsx)
    except (ImportError_, OSError) as e:
        print(f"가져오기 실패: {e}")
        return 1
    finally:
        conn.close()
    print(format_report(rep))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
