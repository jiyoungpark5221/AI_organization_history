"""엑셀 가져오기 시험. 실제 파일(있으면)과 같은 헤더 위치의 소형 가짜 엑셀을 쓴다."""
from pathlib import Path

import openpyxl
import pytest

from app import import_xlsx as imp
from app.collect import collect
from app.db import connect, init_db

REAL = next(Path(__file__).resolve().parent.parent.glob("*Pilot_v03*.xlsx"), None)

ORG_H = ["COMMON_ORG_ID", "DISPLAY_NAME", "PRIMARY_GCC_CODE", "MASTER_STATUS", "PERIOD_STATUS"]
SRC_H = ["SOURCE_ROW_ID", "COMMON_ORG_ID", "SOURCE_SYSTEM", "SOURCE_CODE", "SOURCE_FULL_ORG_NAME",
         "SOURCE_LOWEST_ORG_NAME", "DEPTH", "PARENT_CODE", "TOP_CODE", "SOURCE_START", "SOURCE_END",
         "PREVIOUS_CODE", "SOURCE_STATUS", "RAW_SOURCE_ID", "RAW_SOURCE_LOCATOR"]
REL_H = ["RELATION_ID", "RELATION_LAYER", "FROM_ORG_ID", "RELATION_TYPE", "TO_ORG_ID", "SOURCE_SYSTEM",
         "SOURCE_FIELD", "FROM_SOURCE_CODE", "TO_SOURCE_CODE", "OBSERVED_FROM", "OBSERVED_TO",
         "RELATION_VALID_FROM", "RELATION_VALID_TO", "EVIDENCE", "REVIEW_STATUS", "AI_USE_POLICY",
         "RELATION_DATE_STATUS"]
WORK_H = ["WORK_ID", "COMMON_ORG_ID", "OG_CODE", "REPRESENTATIVE_NAME", "RAW_WORK_TEXT", "RAW_START", "RAW_END",
          "SOURCE_ID", "SOURCE_FILE", "SOURCE_PAGE", "SOURCE_SEQUENCE", "BRM_FUNCTION_STATUS", "WORK_VALID_FROM",
          "WORK_VALID_TO", "WORK_PERIOD_STATUS", "WORK_ADOPTION_STATUS", "LAW_REFERENCE_RAW"]


def row(h, **kw):
    return [kw.get(c) for c in h]


def make_xlsx(path, *, bad_rel=False):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("00_PILOT_SUMMARY")
    ws.append(["요약"])
    ws = wb.create_sheet("01_ORG_MASTER")
    ws.append(["제목"]); ws.append([]); ws.append(ORG_H)
    ws.append(row(ORG_H, COMMON_ORG_ID="ENV-ORG-00001", DISPLAY_NAME="가상부", PRIMARY_GCC_CODE="0100001",
                  MASTER_STATUS="CROSSCHECKED", PERIOD_STATUS="DATE_REVIEW"))
    ws.append(row(ORG_H, COMMON_ORG_ID="ENV-ORG-00002", DISPLAY_NAME="가상부 가상국", PRIMARY_GCC_CODE="0100002",
                  MASTER_STATUS="GCC_ONLY_UNCHECKED", PERIOD_STATUS="GCC_ONLY_UNCHECKED"))
    ws.append(row(ORG_H, COMMON_ORG_ID="ENV-ORG-00003", DISPLAY_NAME="가상처(CAMS만)", MASTER_STATUS="CAMS_ONLY"))
    ws = wb.create_sheet("02_SOURCE_MAP")
    ws.append(["제목"]); ws.append([]); ws.append(SRC_H)
    ws.append(row(SRC_H, SOURCE_ROW_ID="S1", COMMON_ORG_ID="ENV-ORG-00001", SOURCE_SYSTEM="GCC", SOURCE_CODE="0100001",
                  SOURCE_FULL_ORG_NAME="가상부", SOURCE_LOWEST_ORG_NAME="가상부", DEPTH="1", PARENT_CODE="0000000",
                  TOP_CODE="0100001", SOURCE_START="2000-01-01", SOURCE_END=None, PREVIOUS_CODE="NULL",
                  SOURCE_STATUS="GCC_SOURCE", RAW_SOURCE_ID="SRC-X", RAW_SOURCE_LOCATOR="행 2"))
    ws.append(row(SRC_H, SOURCE_ROW_ID="S2", COMMON_ORG_ID="ENV-ORG-00002", SOURCE_SYSTEM="GCC", SOURCE_CODE="0100002",
                  SOURCE_FULL_ORG_NAME="가상부 가상국", SOURCE_LOWEST_ORG_NAME="가상국", DEPTH="2", PARENT_CODE="0100001",
                  TOP_CODE="0100001", SOURCE_START=None, SOURCE_END="2010-05-06", PREVIOUS_CODE="0999999",
                  SOURCE_STATUS="GCC_SOURCE", RAW_SOURCE_ID="SRC-X", RAW_SOURCE_LOCATOR="행 3"))
    ws.append(row(SRC_H, SOURCE_ROW_ID="O1", COMMON_ORG_ID="ENV-ORG-00003", SOURCE_SYSTEM="CAMS_OG", SOURCE_CODE=None,
                  SOURCE_FULL_ORG_NAME="가상처(CAMS만)", SOURCE_START="1990-01-01", SOURCE_END="1994-12-23",
                  PREVIOUS_CODE=None, SOURCE_STATUS="CAMS_ONLY_NO_CODE", RAW_SOURCE_ID="SRC-OG", RAW_SOURCE_LOCATOR="p1"))
    ws = wb.create_sheet("03_ORG_RELATION")
    ws.append(["제목"]); ws.append([]); ws.append(REL_H)
    ws.append(row(REL_H, RELATION_ID="R1", RELATION_LAYER="SOURCE_LINK", FROM_ORG_ID="ENV-ORG-00002",
                  RELATION_TYPE="GCC_PARENT_LINK", TO_ORG_ID="ENV-ORG-00001", SOURCE_SYSTEM="GCC",
                  SOURCE_FIELD="차상위기관코드", FROM_SOURCE_CODE="0100002", TO_SOURCE_CODE="0100001",
                  OBSERVED_FROM="2000-01-01", EVIDENCE="e", REVIEW_STATUS="SOURCE_EXPLICIT",
                  AI_USE_POLICY="CONDITIONAL", RELATION_DATE_STATUS="미확인"))
    ws.append(row(REL_H, RELATION_ID="R2", RELATION_LAYER="CANONICAL_RELATION", FROM_ORG_ID="ENV-ORG-00003",
                  RELATION_TYPE="SUCCEEDS", TO_ORG_ID="ENV-ORG-00001", SOURCE_SYSTEM="GCC+CAMS",
                  SOURCE_FIELD="교차판정", FROM_SOURCE_CODE=None, TO_SOURCE_CODE="0100001",
                  EVIDENCE="e2", REVIEW_STATUS="CONFIRMED_CROSS_SOURCE", AI_USE_POLICY="YES",
                  RELATION_DATE_STATUS="미확인"))
    ws.append(row(REL_H, RELATION_ID="R3", RELATION_LAYER="SOURCE_LINK", FROM_ORG_ID="ENV-ORG-00002",
                  RELATION_TYPE="GCC_PREVIOUS_CODE_LINK", TO_ORG_ID="ENV-ORG-99999" if bad_rel else "ENV-ORG-00001",
                  FROM_SOURCE_CODE="0100002", TO_SOURCE_CODE="0999999"))
    ws = wb.create_sheet("09_ORG_WORK")
    ws.append(["제목"]); ws.append(["설명"]); ws.append([]); ws.append([])
    ws.append(WORK_H)
    ws.append(row(WORK_H, WORK_ID="W1", COMMON_ORG_ID="ENV-ORG-00002", OG_CODE="0100002", REPRESENTATIVE_NAME="가상국",
                  RAW_WORK_TEXT="가상 업무 원문", SOURCE_ID="SRC-OG-1", SOURCE_FILE="a.pdf", SOURCE_PAGE="2",
                  SOURCE_SEQUENCE="1", BRM_FUNCTION_STATUS="미확인", WORK_VALID_FROM="미확인", WORK_VALID_TO="미확인",
                  WORK_PERIOD_STATUS="미확인", WORK_ADOPTION_STATUS="OG 주요업무 원문 반영"))
    ws = wb.create_sheet("14_SOURCE_REGISTER")
    ws.append(["x"])
    wb.save(path)
    return path


@pytest.fixture
def conn(tmp_path):
    c = connect(str(tmp_path / "t.db"))
    init_db(c)
    yield c
    c.close()


@pytest.fixture
def xlsx(tmp_path):
    return make_xlsx(tmp_path / "fake.xlsx")


def n(conn, t):
    return conn.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]


def test_counts_and_report(conn, xlsx):
    rep = imp.run_import(conn, xlsx)
    assert rep["counts"]["orgs_in_file"] == 3 and n(conn, "org") == 3
    assert rep["counts"]["ref_org_record"] == {"GCC": 2, "CAMS_OG": 1}
    assert rep["counts"]["ref_org_work"] == 1 == n(conn, "ref_org_work")
    assert rep["counts"]["ref_org_relation"] == {"SOURCE_LINK": 2, "CANONICAL_RELATION": 1}
    assert rep["source_file_unchanged"] and rep["sha256_before"] == rep["sha256_after"]
    assert "14_SOURCE_REGISTER" in rep["not_imported_sheets"] and "00_PILOT_SUMMARY" in rep["not_imported_sheets"]
    assert "변경 전·후 GCC 자료가 필요" in rep["change_detection_notice"]
    assert "스냅샷이 아니다" in imp.format_report(rep)
    assert n(conn, "import_run") == 1


def test_identifiers_preserved_and_no_og(conn, xlsx):
    imp.run_import(conn, xlsx)
    ids = {(r["id_type"], r["id_value"]) for r in conn.execute("SELECT * FROM org_identifier")}
    assert ("COMMON", "ENV-ORG-00001") in ids and ("GCC", "0100001") in ids  # 선행 0 보존
    assert not any(t == "OG" for t, _ in ids)
    assert not any(t == "GCC" and v == "0999999" for t, v in ids)
    w = conn.execute("SELECT * FROM ref_org_work").fetchone()
    assert w["og_code_raw"] == "0100002"  # OG_CODE 는 원문 보존만
    assert imp.run_import(conn, xlsx)["og_authority_id"]["imported"] == 0


def test_layers_status_and_dates_kept_as_source(conn, xlsx):
    imp.run_import(conn, xlsx)
    rels = {r["relation_id"]: r for r in conn.execute("SELECT * FROM ref_org_relation")}
    assert rels["R1"]["layer"] == "SOURCE_LINK" and rels["R2"]["layer"] == "CANONICAL_RELATION"
    assert rels["R2"]["review_status_src"] == "CONFIRMED_CROSS_SOURCE"
    assert rels["R2"]["relation_valid_from"] is None and rels["R1"]["observed_from"] == "2000-01-01"
    assert rels["R2"]["from_org_id"] is not None and rels["R2"]["origin"] == "imported_reference"
    rec = {r["source_row_id"]: r for r in conn.execute("SELECT * FROM ref_org_record")}
    assert rec["S1"]["previous_code"] is None and rec["S1"]["source_end"] is None  # 'NULL' → NULL
    assert rec["S1"]["master_status_src"] == "CROSSCHECKED"
    assert rec["O1"]["source_system"] == "CAMS_OG" and rec["O1"]["source_code"] is None


def test_no_approval_or_snapshot_data_written(conn, xlsx):
    imp.run_import(conn, xlsx)
    for t in ["confirmed_org_info", "confirmed_relation", "confirmed_task_move", "review_log", "review_case",
              "relation_candidate", "task_move_candidate", "evidence", "correction_log",
              "collection_run", "gcc_record", "change_event"]:
        assert n(conn, t) == 0, t


def test_first_collect_after_import_is_still_baseline(conn, xlsx):
    imp.run_import(conn, xlsx)

    class Gcc:
        def fetch(self, day):
            return {"complete": True, "records": [
                {"gcc_code": "0100001", "org_name": "가상부", "parent_code": None, "prev_org_code": None,
                 "created_date": "2000-01-01", "abolished_date": None}]}

    res = collect(conn, Gcc(), "2026-10-01")
    assert res["is_baseline"] and res["change_count"] == 0 and n(conn, "review_case") == 0


def test_idempotent_reimport(conn, xlsx):
    imp.run_import(conn, xlsx)
    tables = ["org", "org_identifier", "ref_org_record", "ref_org_work", "ref_org_relation"]
    before = {t: n(conn, t) for t in tables}
    rep = imp.run_import(conn, xlsx)
    assert {t: n(conn, t) for t in tables} == before
    assert rep["counts"].get("orgs_created", 0) == 0 and rep["counts"].get("gcc_identifiers_added", 0) == 0
    assert n(conn, "import_run") == 2


def test_unlinked_codes_and_unknown_dates(conn, xlsx):
    rep = imp.run_import(conn, xlsx)
    u, d = rep["unlinked"], rep["unknown_dates"]
    assert [x["source_row_id"] for x in u["rows_without_code"]] == ["O1"]
    assert u["previous_code_not_in_gcc"]["sample"] == ["0999999"]
    assert u["parent_code_not_in_gcc"]["rows"] == 0  # '0000000' 은 상위 없음
    assert u["relation_to_code_not_in_gcc"]["sample"] == ["0999999"]
    assert u["relation_org_unresolved"] == 0
    assert d["source_start_blank"] == 1 and d["source_end_blank"] == 1
    assert d["relation_valid_from_blank"] == d["relation_total"] == 3
    assert d["work_period_status_unknown"] == 1 and d["work_valid_from_unknown"] == 1


def test_unresolved_relation_org_counted(conn, tmp_path):
    rep = imp.run_import(conn, make_xlsx(tmp_path / "b.xlsx", bad_rel=True))
    assert rep["unlinked"]["relation_org_unresolved"] == 1
    r = conn.execute("SELECT to_org_id FROM ref_org_relation WHERE relation_id='R3'").fetchone()
    assert r["to_org_id"] is None


def test_reuses_existing_org_with_same_gcc_code(conn, xlsx):
    oid = conn.execute("INSERT INTO org(org_name) VALUES('기존 기관')").lastrowid
    conn.execute("INSERT INTO org_identifier VALUES(?, 'GCC', '0100001')", (oid,))
    rep = imp.run_import(conn, xlsx)
    assert rep["counts"]["orgs_matched_existing"] == 1 and n(conn, "org") == 3
    assert conn.execute("SELECT org_name FROM org WHERE org_id=?", (oid,)).fetchone()[0] == "기존 기관"


def test_failure_rolls_back_everything(conn, xlsx, monkeypatch):
    def boom(*a, **k):
        raise RuntimeError("주입된 실패")

    monkeypatch.setattr(imp, "_write_relations", boom)
    with pytest.raises(RuntimeError):
        imp.run_import(conn, xlsx)
    for t in ["org", "org_identifier", "ref_org_record", "ref_org_work", "ref_org_relation", "import_run"]:
        assert n(conn, t) == 0, t


def test_source_file_not_modified(conn, xlsx):
    before = imp.sha256_of(xlsx)
    imp.run_import(conn, xlsx)
    assert imp.sha256_of(xlsx) == before


def test_old_db_without_common_type_is_migrated(tmp_path, xlsx):
    c = connect(str(tmp_path / "old.db"))
    init_db(c)
    c.execute("DROP TABLE org_identifier")
    c.execute("CREATE TABLE org_identifier (org_id INTEGER NOT NULL REFERENCES org(org_id), "
              "id_type TEXT NOT NULL CHECK (id_type IN ('GCC','OG')), id_value TEXT NOT NULL, "
              "PRIMARY KEY (id_type, id_value))")
    oid = c.execute("INSERT INTO org(org_name) VALUES('x')").lastrowid
    c.execute("INSERT INTO org_identifier VALUES(?, 'GCC', '5555555')", (oid,))
    imp.run_import(c, xlsx)
    assert c.execute("SELECT COUNT(*) FROM org_identifier WHERE id_value='5555555'").fetchone()[0] == 1
    assert c.execute("SELECT COUNT(*) FROM org_identifier WHERE id_type='COMMON'").fetchone()[0] == 3
    c.close()


@pytest.mark.skipif(REAL is None, reason="실제 엑셀 파일 없음")
def test_real_file(conn):
    before = imp.sha256_of(REAL)
    rep = imp.run_import(conn, REAL)
    assert imp.sha256_of(REAL) == before and rep["source_file_unchanged"]
    assert n(conn, "ref_org_work") == rep["counts"]["ref_org_work"] > 0
    assert sum(rep["counts"]["ref_org_relation"].values()) == n(conn, "ref_org_relation")
    assert set(rep["counts"]["ref_org_relation"]) == {"SOURCE_LINK", "CANONICAL_RELATION"}
    assert conn.execute("SELECT COUNT(*) FROM org_identifier WHERE id_type='OG'").fetchone()[0] == 0
    assert n(conn, "confirmed_relation") == n(conn, "collection_run") == n(conn, "review_log") == 0
    w = rep["og_code_equals_gcc_code"]
    assert w["work_rows_with_og_code"] == w["same_as_gcc_code_of_same_org"]
